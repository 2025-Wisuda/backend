"""
Custom Excel export format with auto-fit columns and professional styling.

This module provides an enhanced XLSX export format for tabular data, automatically adjusting column widths and applying professional header styling for improved readability and internationalization support.

Key Features:
- Automatic column width adjustment based on content (auto-fit)
- Professional header styling (bold, centered)
- Padding and max-width for better display
- Handles international (non-ASCII) characters
- Fallback to standard XLSX export on error
"""

import logging
from io import BytesIO

from import_export.formats.base_formats import XLSX
from openpyxl import load_workbook
from openpyxl.styles import Alignment, Font
from openpyxl.utils import get_column_letter

logger = logging.getLogger(__name__)


class AutoFitXLSX(XLSX):
    """
    Enhanced XLSX format with auto-fitting columns and professional styling.

    Extends the standard XLSX format to provide:
    - Automatic column width adjustment based on content
    - Professional header styling with bold, centered fonts
    - Improved readability with padding and max-width
    - Support for international characters
    - Fallback to standard XLSX export if an error occurs
    """

    def export_data(self, dataset, **kwargs) -> bytes:
        """
        Export dataset to XLSX format with auto-fit columns and styling.

        Args:
            dataset: The dataset to export (tabular data).
            **kwargs: Additional export parameters.

        Returns:
            bytes: The formatted XLSX file content.
        """
        try:
            raw_data = dataset.export("xlsx", **kwargs)

            wb = load_workbook(filename=BytesIO(raw_data))
            ws = wb.active

            if ws.max_row > 0:
                for cell in ws[1]:
                    cell.font = Font(bold=True)
                    cell.alignment = Alignment(horizontal="center", vertical="center")

            for col_idx, col_cells in enumerate(ws.columns, 1):
                max_length = 0
                column_letter = get_column_letter(col_idx)

                for cell in col_cells:
                    if cell.value is not None:
                        cell_value = str(cell.value)
                        char_length = len(cell_value)

                        if any(ord(c) > 127 for c in cell_value):
                            char_length = int(char_length * 1.3)

                        max_length = max(max_length, char_length)

                min_width = 10
                padding = 4
                adjusted_width = max(min_width, max_length + padding)

                max_width = 60
                adjusted_width = min(adjusted_width, max_width)

                ws.column_dimensions[column_letter].width = adjusted_width

                logger.debug(
                    f"Column {column_letter}: max_length={max_length}, adjusted_width={adjusted_width}"
                )

            output = BytesIO()
            wb.save(output)
            return output.getvalue()

        except Exception as e:
            logger.error(f"Error in AutoFitXLSX export: {str(e)}")
            return dataset.export("xlsx", **kwargs)

    def get_title(self) -> str:
        """
        Get the display title for this export format.

        Returns:
            str: The display title for this format (for UI dropdowns, etc).
        """
        return "XLSX (Auto-fit)"


CustomXLSXFormat = AutoFitXLSX
